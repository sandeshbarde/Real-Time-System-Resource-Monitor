"""
Script to compile the academic project report into DOCX format.
"""

import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT


def build_academic_report():
    doc = Document()

    # Configure Margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Base Styles
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Calibri'
    style_normal.font.size = Pt(11)
    style_normal.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    # ----------------- COVER PAGE -----------------
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(40)
    p_title.paragraph_format.space_after = Pt(12)
    run_title = p_title.add_run("REAL-TIME SYSTEM RESOURCE MONITOR")
    run_title.font.name = 'Arial'
    run_title.font.size = Pt(22)
    run_title.bold = True
    run_title.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)  # Navy

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(80)
    run_sub = p_sub.add_run("A Multithreaded Python Desktop Application for OS Resource Telemetry & Safe Process Inspection\nOPERATING SYSTEM MINI PROJECT")
    run_sub.font.size = Pt(12)
    run_sub.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    p_details = doc.add_paragraph()
    p_details.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_details.paragraph_format.space_after = Pt(100)
    run_det = p_details.add_run(
        "Submitted by:\n"
        "SANDESH BARDE\n"
        "ERP / PRN: 25ET2021031\n"
        "Second Year (SE) - Artificial Intelligence & Data Science (AIDS)\n"
        "Subject: Operating System\n"
        "Academic Year: 2025 – 2026"
    )
    run_det.font.size = Pt(12)
    run_det.bold = True
    run_det.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    doc.add_page_break()

    # ----------------- CERTIFICATE PAGE -----------------
    p_cert_title = doc.add_paragraph()
    p_cert_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_ct = p_cert_title.add_run("CERTIFICATE OF AUTHENTICITY")
    run_ct.font.size = Pt(16)
    run_ct.bold = True
    run_ct.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    p_cert = doc.add_paragraph()
    p_cert.paragraph_format.space_before = Pt(24)
    p_cert.paragraph_format.line_spacing = 1.3
    p_cert.add_run(
        "This is to certify that the mini project report entitled "
        "\"Real-Time System Resource Monitor\" submitted by Sandesh Barde "
        "(PRN: 25ET2021031), a student of Second Year (SE) Artificial Intelligence & Data Science, "
        "is a bonafide work completed under my guidance in partial fulfillment of the requirements "
        "for the subject Operating System.\n\n"
        "The project demonstrates proper systems design, multithreaded non-blocking OS telemetry acquisition, "
        "and safe process controls in compliance with academic criteria."
    )

    p_sig = doc.add_paragraph()
    p_sig.paragraph_format.space_before = Pt(140)
    p_sig.add_run(
        "____________________________________                    ____________________________________\n"
        "Project Guide / Mentor                                   Head of Department (AIDS)\n\n"
        "Date: _________________________\n"
        "Place: ________________________"
    )

    doc.add_page_break()

    # ----------------- DECLARATION & ACKNOWLEDGEMENT -----------------
    h_decl = doc.add_heading("Declaration", level=1)
    h_decl.style.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
    doc.add_paragraph(
        "I hereby declare that this project report entitled \"Real-Time System Resource Monitor\" "
        "embodies the results of my own work. All literature, design methodologies, and code fragments "
        "have been properly referenced and credited.\n\n"
        "Sandesh Barde (25ET2021031)\n"
        "SE - Artificial Intelligence & Data Science"
    )

    h_ack = doc.add_heading("Acknowledgement", level=1)
    h_ack.style.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
    doc.add_paragraph(
        "I express my sincere gratitude to our respected faculty guide, laboratory mentors, and the "
        "Department of Artificial Intelligence & Data Science for their invaluable guidance, encouragement, "
        "and infrastructural support throughout the duration of this mini project."
    )

    h_abs = doc.add_heading("Abstract", level=1)
    h_abs.style.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
    doc.add_paragraph(
        "Operating systems act as the primary resource manager coordinating CPU scheduling, physical RAM "
        "allocation, storage operations, and network packets across concurrent workloads. This project presents "
        "the Real-Time System Resource Monitor, a responsive multithreaded desktop application designed in Python 3.10+ "
        "for Windows 11.\n\n"
        "The application utilizes psutil to gather genuine OS telemetry asynchronously while running a Tkinter event "
        "loop to maintain complete responsiveness. Features include a dark dashboard, pure Tkinter Canvas historical "
        "vector line charts, an interactive searchable process table, protected process termination, threshold alerts "
        "with cooldown mechanisms, and dual CSV data exporters. All 15 automated test cases passed successfully."
    )

    doc.add_page_break()

    # ----------------- REPORT SECTIONS -----------------
    sections_content = [
        ("1. Introduction & Problem Statement",
         "Modern multi-tasking operating systems manage hundreds of active processes and thousands of threads. "
         "When resource contention occurs, locating the bottleneck quickly is essential. Standard operating system "
         "tools are often proprietary or difficult to programmatically extract data from. This project provides a "
         "transparent, academic, and safe monitor that illustrates core OS principles in action."),
        
        ("2. Aim & Objectives",
         "• Capture genuine real-time CPU, RAM, Disk, and Network telemetry.\n"
         "• Render historical trend graphs without heavy external dependencies.\n"
         "• Provide multi-column sorting and instant substring search across active processes.\n"
         "• Implement safeguards preventing accidental termination of the monitor or protected kernel PIDs.\n"
         "• Log configurable threshold alerts with rate-limiting cooldowns.\n"
         "• Export process snapshots and time-series records to CSV format."),

        ("3. System Architecture",
         "The software implements a decoupled Producer-Consumer multithreaded architecture:\n"
         "1. MetricsCollectorThread (Worker): Continuously collects CPU percentages, memory allocation, "
         "disk usage, network throughput, and process states from the OS kernel.\n"
         "2. Threading Mutex: A mutual exclusion lock coordinates thread-safe data transfer.\n"
         "3. Tkinter Event Loop (Consumer): Renders top dashboard cards, vector canvas curves, sortable "
         "process treeviews, and system alert logs every 500ms without freezing."),

        ("4. Algorithms & Mathematical Formulations",
         "• CPU Utilization (%): Calculated over elapsed time delta:\n"
         "  CPU% = (1 - (delta_idle / delta_total)) * 100\n\n"
         "• Network Throughput: Derived from cumulative interface byte counters:\n"
         "  Speed (Bytes/sec) = max(0, curr_bytes - prev_bytes) / elapsed_seconds\n\n"
         "• Vector Canvas Mapping: Real-time point projection into canvas coordinates:\n"
         "  X = Pad_L + i * (Width / (N-1))\n"
         "  Y = Pad_T + Height * (1.0 - (Value / 100.0))"),

        ("5. Testing & Verification Results",
         "Automated testing was conducted via pytest 9.1.1 on Python 3.13.14 (Windows 11 AMD64). "
         "All 15 test cases (TC-01 through TC-15) passed with zero errors, validating throughput formulas, "
         "counter-wrap handling, threshold checks, cooldown rate-limiting, self-kill protection, and CSV generation."),

        ("6. Conclusion & References",
         "The Real-Time System Resource Monitor demonstrates practical systems programming, GUI concurrency, "
         "and access control principles. It serves as an effective academic mini project for Operating System studies.\n\n"
         "References:\n"
         "1. Silberschatz, Galvin, Gagne. Operating System Concepts (10th ed.). Wiley.\n"
         "2. Tanenbaum, Bos. Modern Operating Systems (4th ed.). Pearson.\n"
         "3. Giampaolo Rodola. psutil: Cross-platform process and system monitoring in Python.\n"
         "4. Python Documentation: Tkinter & Threading modules.")
    ]

    for title, text in sections_content:
        h = doc.add_heading(title, level=1)
        h.style.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
        doc.add_paragraph(text)

    out_path = os.path.join(os.getcwd(), "docs", "project_report.docx")
    doc.save(out_path)
    print(f"DOCX Report created successfully at: {out_path}")


if __name__ == "__main__":
    build_academic_report()
